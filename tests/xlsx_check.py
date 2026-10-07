#!/usr/bin/env python3
"""
Independent validation of the .xlsx files written by app/lib/xlsx.php.

Reads them with openpyxl (a different implementation from ours) and, if LibreOffice
is installed, also opens them headless in LibreOffice Calc and checks the values as
they are displayed (number formats applied).

Usage: python3 -I tests/xlsx_check.py tests/output/xlsx_expect.json
       (tests/xlsx.php writes the expectations file and runs this script)
Exit: 0 = pass, 1 = failures, 3 = openpyxl not installed (skip)
Env:  XLSX_CHECK_NO_SOFFICE=1 skips the LibreOffice part.
"""
import csv
import datetime
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import time
import zipfile
import xml.etree.ElementTree as ET

try:
    import openpyxl
    from openpyxl.utils import get_column_letter
except ImportError:
    print("SKIP  openpyxl is not installed (pip install openpyxl)")
    sys.exit(3)

PASS = 0
FAILS = []


def check(label, ok, detail=""):
    global PASS
    if ok:
        PASS += 1
        print("  PASS  " + label)
    else:
        FAILS.append(label)
        print("  FAIL  " + label + ("  -> " + detail if detail else ""))


def excel_text(s):
    """Text as Excel displays it: _xHHHH_ escapes decoded (openpyxl leaves them as-is)."""
    return re.sub(r"_x([0-9A-Fa-f]{4})_", lambda m: chr(int(m.group(1), 16)), s)


def same_value(kind, expected, cell):
    """Compare an openpyxl cell with an expectation [kind, value, number_format]."""
    v = cell.value
    if kind is None:
        return v is None, "expected empty, got %r" % (v,)
    if kind == "n":
        ok = cell.data_type == "n" and isinstance(v, (int, float)) and not isinstance(v, bool) and float(v) == float(expected)
        return ok, "expected number %s, got %r (%s)" % (expected, v, cell.data_type)
    if kind == "d":
        ok = cell.data_type == "d" and isinstance(v, datetime.datetime) and v == datetime.datetime.fromisoformat(expected)
        return ok, "expected datetime %s, got %r (%s)" % (expected, v, cell.data_type)
    ok = cell.data_type == "s" and isinstance(v, str) and excel_text(v) == expected
    return ok, "expected text %r, got %r (%s)" % (expected[:60], v if not isinstance(v, str) else v[:60], cell.data_type)


def check_package(path):
    with zipfile.ZipFile(path) as z:
        check("zipfile CRC check: " + os.path.basename(path), z.testzip() is None)
        bad = []
        for name in z.namelist():
            if name.endswith(".xml") or name.endswith(".rels"):
                try:
                    ET.fromstring(z.read(name))
                except ET.ParseError as e:
                    bad.append("%s: %s" % (name, e))
        check("every XML part parses (ElementTree)", not bad, "; ".join(bad))


def check_sample(exp):
    path = exp["path"]
    check_package(path)
    wb = openpyxl.load_workbook(path)  # not data_only: formulas would show as data_type 'f'
    check("one sheet named %r" % exp["sheet"], wb.sheetnames == [exp["sheet"]], repr(wb.sheetnames))
    ws = wb[wb.sheetnames[0]]
    check("sheet is right-to-left", ws.sheet_view.rightToLeft is True)
    check("freeze pane below the header (%s)" % exp["freeze"], ws.freeze_panes == exp["freeze"], repr(ws.freeze_panes))
    check("autofilter range %s" % exp["autofilter"], ws.auto_filter.ref == exp["autofilter"], repr(ws.auto_filter.ref))
    merged = sorted(str(r) for r in ws.merged_cells.ranges)
    check("merged title rows", merged == sorted(exp["merged"]), repr(merged))
    check("document title", wb.properties.title == exp["title"], repr(wb.properties.title))
    check("document creator = company name", wb.properties.creator == exp["creator"], repr(wb.properties.creator))
    titles = str(ws.print_title_rows or "")
    check("header row repeats when printing", titles.replace("$", "") in ("4:4",), repr(titles))

    hr = exp["header_row"]
    labels = [ws.cell(row=hr, column=i + 1).value for i in range(len(exp["header"]))]
    check("header labels (Arabic, exact)", labels == exp["header"], repr(labels))
    head = [ws.cell(row=hr, column=i + 1) for i in range(len(exp["header"]))]
    check("header: bold, #F5F5F4 fill, thin bottom border",
          all(c.font.b and c.fill.fgColor.rgb == "FFF5F5F4" and c.border.bottom.style == "thin" for c in head))
    tr = exp["total_row"]
    tot = [ws.cell(row=tr, column=i + 1) for i in range(len(exp["header"]))]
    check("totals row: bold with a thin top border", all(c.font.b and c.border.top.style == "thin" for c in tot))
    check("title: bold 16pt", ws["A1"].font.b and float(ws["A1"].font.sz) == 16.0)

    bad = []
    for ref, e in exp["cells"].items():
        cell = ws[ref]
        ok, detail = same_value(e[0], e[1] if len(e) > 1 else None, cell)
        if ok and e[0] is not None and cell.number_format != e[2]:
            ok, detail = False, "format %r != %r" % (cell.number_format, e[2])
        if not ok:
            bad.append("%s: %s" % (ref, detail))
    check("%d cells: value, type and number format" % len(exp["cells"]), not bad, "; ".join(bad))

    formulas = [c.coordinate for row in ws.iter_rows() for c in row if c.data_type == "f"]
    check("no formulas anywhere (formula-looking text stays text)", not formulas, repr(formulas[:5]))
    widths = {k: d.width for k, d in ws.column_dimensions.items()}
    check("column widths set and capped at 60", len(widths) == len(exp["header"]) and all(0 < w <= 60 for w in widths.values()), repr(widths))
    check("long and multi-line text wraps", all(ws[r].alignment.wrap_text for r in exp["wrapped"]))
    check("start = right, end = left (RTL sheet)", ws["A5"].alignment.horizontal == "right" and ws["D5"].alignment.horizontal == "left")


def check_big(exp):
    path = exp["path"]
    check_package(path)
    t = time.time()
    wb = openpyxl.load_workbook(path, read_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(min_row=5, values_only=True))
    secs = time.time() - t
    print("  INFO  openpyxl read %d rows in %.2f s" % (len(rows), secs))
    n = exp["rows"]
    check("rows read back: %d data + totals" % n, len(rows) == n + 1, str(len(rows)))
    data, totals = rows[:n], rows[n] if len(rows) > n else ()

    def norm(row):
        out = []
        for v in row:
            if isinstance(v, datetime.datetime):
                out.append(v.isoformat())
            else:
                out.append(v)
        return out

    def expect_row(e):
        return [datetime.datetime.fromisoformat(v).isoformat() if isinstance(v, str) and re.match(r"^\d{4}-\d\d-\d\dT", v)
                else (float(v) if isinstance(v, str) and re.match(r"^-?\d+(\.\d+)?$", v) else v) for v in e]

    def eq(row, e):
        got = norm(row)
        want = expect_row(e)
        return len(got) == len(want) and all(
            (isinstance(g, (int, float)) and isinstance(w, (int, float)) and float(g) == float(w)) or g == w
            for g, w in zip(got, want))

    check("first data row", eq(data[0], exp["first"]), repr(norm(data[0])))
    check("last data row", eq(data[-1], exp["last"]), repr(norm(data[-1])))
    wrong = 0
    for row in data:
        for v, kind in zip(row, exp["types"]):
            if kind == "s?" and v is None:
                continue
            ok = (kind.startswith("s") and isinstance(v, str)) or \
                 (kind == "n" and isinstance(v, (int, float)) and not isinstance(v, bool)) or \
                 (kind == "d" and isinstance(v, datetime.datetime))
            if not ok:
                wrong += 1
    check("every data cell has the right type (text/number/datetime)", wrong == 0, "%d wrong" % wrong)
    tt = exp["totals"]
    check("totals row", eq(totals, [v if v is not None else None for v in tt]), repr(norm(totals)))
    wb.close()


def check_libreoffice(exp):
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice or os.environ.get("XLSX_CHECK_NO_SOFFICE"):
        print("  SKIP  LibreOffice not available")
        return
    work = tempfile.mkdtemp(prefix="xlsx-check-")
    try:
        profile = "file://" + os.path.join(work, "profile")
        # 44=comma, 34=quote, 76=UTF-8, 9th token true = save cell content as shown (formatted)
        filt = "csv:Text - txt - csv (StarCalc):44,34,76,1,,0,false,true,true"
        env = dict(os.environ, LC_ALL="C.UTF-8", LANG="C.UTF-8")
        t = time.time()
        r = subprocess.run([soffice, "-env:UserInstallation=" + profile, "--headless", "--convert-to", filt,
                            "--outdir", work, exp["path"]], capture_output=True, text=True, timeout=180, env=env)
        out = os.path.join(work, os.path.splitext(os.path.basename(exp["path"]))[0] + ".csv")
        check("LibreOffice opens and converts the file", r.returncode == 0 and os.path.exists(out), (r.stderr or r.stdout)[-300:])
        if not os.path.exists(out):
            return
        print("  INFO  LibreOffice conversion took %.1f s" % (time.time() - t))
        with open(out, encoding="utf-8-sig", newline="") as f:
            grid = list(csv.reader(f))
        bad = []
        for ref, want in exp["shown"].items():
            m = re.match(r"^([A-Z]+)(\d+)$", ref)
            col = sum((ord(ch) - 64) * 26 ** i for i, ch in enumerate(reversed(m.group(1)))) - 1
            row = int(m.group(2)) - 1
            got = grid[row][col] if row < len(grid) and col < len(grid[row]) else None
            if got != want:
                bad.append("%s: %r != %r" % (ref, got, want))
        check("LibreOffice displays %d cells with the expected formatting" % len(exp["shown"]), not bad, "; ".join(bad))
    finally:
        shutil.rmtree(work, ignore_errors=True)


def main():
    exp = json.load(open(sys.argv[1], encoding="utf-8"))
    print("== openpyxl %s: sample workbook" % openpyxl.__version__)
    check_sample(exp["sample"])
    print("== openpyxl: %d-row workbook" % exp["big"]["rows"])
    check_big(exp["big"])
    print("== LibreOffice (optional)")
    check_libreoffice(exp["sample"])
    print("Result: %d passed, %d failed" % (PASS, len(FAILS)))
    for f in FAILS:
        print("  - " + f)
    sys.exit(1 if FAILS else 0)


if __name__ == "__main__":
    main()
